#!/usr/bin/env python3
"""Collect explicit S1/S2/S3 smoke assertions from the isolated cloud database.

The caller writes ``--output`` on cloud_server with mode 0600. API checks are
loopback-only. Temporary S2 snapshot/export and S3 binary-image probes are
removed before returning; no Trail, production, or external notification write
is performed.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import os
import re
import stat
import subprocess
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import openpyxl
from PIL import Image

APP_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_ROOT))

from app.db import Database  # noqa: E402
from app.db_parts.snapshots import _snapshot_membership_sha  # noqa: E402

POSTGRES_MIGRATIONS_DIR = APP_ROOT / "migrations" / "postgres"
SAFE_DATABASE_RE = re.compile(r"^manual_s3_smoke(?:_[a-z0-9][a-z0-9_-]*)?$", re.IGNORECASE)
SAFE_HOSTS = {"", "127.0.0.1", "::1", "localhost"}
REQUIRED_MODEL_STATUSES = {"pending", "in_progress", "completed", "blocked_by_label"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database-url-file", required=True)
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--base-url", default="http://127.0.0.1:8786")
    parser.add_argument("--output", required=True)
    parser.add_argument("--backfill-dir", default="")
    parser.add_argument("--performance-report", default="")
    parser.add_argument("--review-attachments-dir", default="")
    parser.add_argument("--pytest-log", default="")
    return parser.parse_args()


def read_url_file(path: Path) -> str:
    descriptor = os.open(path, os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0))
    try:
        metadata = os.fstat(descriptor)
        if not stat.S_ISREG(metadata.st_mode) or metadata.st_uid != os.getuid():
            raise RuntimeError("smoke URL file ownership/type check failed")
        if stat.S_IMODE(metadata.st_mode) != 0o600:
            raise RuntimeError("smoke URL file must be 0600")
        value = os.read(descriptor, 16 * 1024 + 1)
    finally:
        os.close(descriptor)
    if len(value) > 16 * 1024:
        raise RuntimeError("smoke URL file is too large")
    result = value.decode("utf-8").strip()
    if not result:
        raise RuntimeError("smoke URL file is empty")
    return result


def api_request(base_url: str, method: str, path: str, body: dict[str, Any] | None = None) -> tuple[int, dict[str, str], bytes]:
    payload = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    headers = {} if payload is None else {"Content-Type": "application/json"}
    request = urllib.request.Request(base_url + path, data=payload, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            return response.status, dict(response.headers), response.read()
    except urllib.error.HTTPError as error:
        return error.code, dict(error.headers), error.read()


def api_json(base_url: str, path: str) -> dict[str, Any]:
    status, _headers, payload = api_request(base_url, "GET", path)
    if not 200 <= status < 300:
        raise RuntimeError(f"API GET returned HTTP {status}")
    return json.loads(payload.decode("utf-8")) if payload else {}


def api_multipart_image(
    base_url: str, path: str, fields: dict[str, str], filename: str, image_bytes: bytes
) -> tuple[int, dict[str, str], bytes]:
    boundary = "----manual-s3-smoke-attachment"
    chunks: list[bytes] = []
    for name, value in fields.items():
        chunks.extend(
            [
                f"--{boundary}\r\n".encode(),
                f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                value.encode("utf-8"),
                b"\r\n",
            ]
        )
    chunks.extend(
        [
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="attachments"; filename="{filename}"\r\n'.encode(),
            b"Content-Type: image/png\r\n\r\n",
            image_bytes,
            b"\r\n",
            f"--{boundary}--\r\n".encode(),
        ]
    )
    request = urllib.request.Request(
        base_url + path,
        data=b"".join(chunks),
        headers={
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "X-RA-Triage-Request": "review-v1",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            return response.status, dict(response.headers), response.read()
    except urllib.error.HTTPError as error:
        return error.code, dict(error.headers), error.read()


def _binary_attachment_probe(
    database: Database,
    manifest: dict[str, Any],
    base_url: str,
    manifest_path: Path,
    attachment_root: Path,
) -> dict[str, Any]:
    smoke_root = manifest_path.parent.resolve()
    expected_root = (smoke_root / "data" / "review_attachments").resolve()
    root = attachment_root.expanduser().resolve()
    if root != expected_root:
        raise RuntimeError("attachment probe root must be the isolated smoke data directory")
    issue_ids, _scope_map = _selected_ids(manifest)
    if not issue_ids:
        raise AssertionError("smoke manifest contains no selected Issue")
    issue_id = issue_ids[0]
    run_id = str(manifest["source_run_ids"][0])
    reviewer = "manual_s3_attachment_smoke_probe"
    with database.connect() as connection:
        exists = int(connection.execute(
            "SELECT COUNT(*) AS n FROM model_review_revisions WHERE issue_id=? AND model_run_id=? AND reviewer=?",
            (issue_id, run_id, reviewer),
        ).fetchone()["n"])
    if exists:
        raise RuntimeError("attachment probe fixture already exists; refusing to overwrite it")

    revision_id = 0
    stored_names: list[str] = []
    try:
        image_buffer = io.BytesIO()
        Image.new("RGB", (1, 1), (48, 96, 160)).save(image_buffer, format="PNG", optimize=True)
        png_bytes = image_buffer.getvalue()
        body = {
            "model_run_id": run_id,
            "model_review_status": "in_progress",
            "expected_previous_annotation_id": None,
            "note": "temporary binary attachment smoke probe",
            "author": reviewer,
        }
        issue_path = urllib.parse.quote(issue_id, safe="")
        status, _headers, response_bytes = api_multipart_image(
            base_url,
            f"/api/cases/{issue_path}/annotations-with-attachments",
            {"payload": json.dumps(body, ensure_ascii=False)},
            "manual-s3-smoke.png",
            png_bytes,
        )
        if status != 200:
            raise AssertionError(f"binary attachment upload returned HTTP {status}")
        response = json.loads(response_bytes.decode("utf-8"))
        public_attachments = (response.get("annotation") or {}).get("attachments") or []
        if len(public_attachments) != 1:
            raise AssertionError("binary attachment upload did not return one attachment")
        attachment_id = str(public_attachments[0].get("id") or "")
        if not attachment_id or public_attachments[0].get("media_type") != "image/png":
            raise AssertionError("binary attachment metadata is incomplete")

        with database.connect() as connection:
            rows = connection.execute(
                """
                SELECT revision.id AS revision_id, attachment.stored_name,
                       attachment.sha256, attachment.width, attachment.height
                FROM model_review_revisions revision
                JOIN model_review_attachments attachment ON attachment.revision_id=revision.id
                WHERE revision.issue_id=? AND revision.model_run_id=? AND revision.reviewer=?
                  AND attachment.id=?
                """,
                (issue_id, run_id, reviewer, attachment_id),
            ).fetchall()
        if len(rows) != 1:
            raise AssertionError("uploaded attachment is not linked to its Model Review revision")
        revision_id = int(rows[0]["revision_id"])
        stored_names.append(str(rows[0]["stored_name"]))
        if int(rows[0]["width"]) != 1 or int(rows[0]["height"]) != 1:
            raise AssertionError("uploaded binary image dimensions changed unexpectedly")
        file_path = (root / stored_names[0]).resolve()
        if root not in file_path.parents or not file_path.is_file():
            raise AssertionError("uploaded image is not stored beneath the smoke attachment root")
        download_status, download_headers, download_bytes = api_request(
            base_url, "GET", f"/api/review-attachments/{urllib.parse.quote(attachment_id, safe='')}"
        )
        content_type = next(
            (str(value) for key, value in download_headers.items() if key.lower() == "content-type"),
            "",
        )
        if download_status != 200 or not content_type.startswith("image/png"):
            raise AssertionError("uploaded binary image cannot be read back through its API")
        if hashlib.sha256(download_bytes).hexdigest() != str(rows[0]["sha256"]):
            raise AssertionError("uploaded binary image read-back hash differs from its stored metadata")
        return {"binary_upload_and_readback": True, "revision_attachment_linked": True, "temporary_artifacts_removed": True}
    finally:
        with database.connect() as connection:
            revisions = [int(row["id"]) for row in connection.execute(
                "SELECT id FROM model_review_revisions WHERE issue_id=? AND model_run_id=? AND reviewer=?",
                (issue_id, run_id, reviewer),
            ).fetchall()]
            if revision_id and revision_id not in revisions:
                revisions.append(revision_id)
            for cleanup_revision_id in revisions:
                attachments = connection.execute(
                    "SELECT stored_name FROM model_review_attachments WHERE revision_id=?",
                    (cleanup_revision_id,),
                ).fetchall()
                stored_names.extend(str(row["stored_name"]) for row in attachments)
                connection.execute("DELETE FROM model_review_heads WHERE revision_id=?", (cleanup_revision_id,))
                connection.execute("DELETE FROM model_review_attachments WHERE revision_id=?", (cleanup_revision_id,))
                connection.execute("DELETE FROM model_review_revisions WHERE id=?", (cleanup_revision_id,))
        for stored_name in set(stored_names):
            file_path = (root / stored_name).resolve()
            if root in file_path.parents:
                file_path.unlink(missing_ok=True)
        with database.connect() as connection:
            leftover = int(connection.execute(
                "SELECT COUNT(*) AS n FROM model_review_revisions WHERE issue_id=? AND model_run_id=? AND reviewer=?",
                (issue_id, run_id, reviewer),
            ).fetchone()["n"])
        if leftover:
            raise RuntimeError("binary attachment smoke probe database cleanup failed")


def _load_smoke_builder() -> Any:
    path = APP_ROOT / "scripts" / "build_manual_s3_smoke.py"
    spec = importlib.util.spec_from_file_location("manual_s3_smoke_builder", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("smoke builder module unavailable")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _read_check(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def record_check(checks: list[dict[str, Any]], name: str, fn: Callable[[], dict[str, Any]]) -> None:
    try:
        details = fn()
        checks.append({"name": name, "status": "PASS", "assertions": details})
    except Exception as exc:
        message = str(exc)
        message = re.sub(r"\bcn[a-z0-9_-]{6,}\b", "<issue-id>", message, flags=re.IGNORECASE)
        message = re.sub(r"(?i)postgres(?:ql)?://[^\s\"'<>]+", "<database-url>", message)
        message = re.sub(r"(?i)(password|token|secret)=\S+", r"\1=<redacted>", message)
        checks.append({
            "name": name,
            "status": "FAIL",
            "error_type": type(exc).__name__,
            "error_summary": message[:240],
        })


def _selected_ids(manifest: dict[str, Any]) -> tuple[list[str], dict[str, str]]:
    ids: list[str] = []
    scope_by_id: dict[str, str] = {}
    for scope in manifest["scopes"]:
        scope_ids = list(dict.fromkeys(scope["hash_sample_issue_ids"] + scope["pinned_addition_issue_ids"]))
        ids.extend(scope_ids)
        for issue_id in scope_ids:
            scope_by_id[str(issue_id)] = str(scope["scope"])
    return list(dict.fromkeys(ids)), scope_by_id


def _api_pair(database: Database, manifest: dict[str, Any], base_url: str) -> dict[str, Any]:
    runs = list(manifest["source_run_ids"])
    if len(runs) != 2:
        raise AssertionError("the smoke manifest does not carry two pinned Runs")
    selected, scope_by_issue = _selected_ids(manifest)
    author = "manual_s3_smoke_gallery_probe"
    with database.connect() as connection:
        rows = connection.execute(
            """
            SELECT revision.issue_id, head.reviewer, head.model_run_id, revision.status,
                   revision.reason, revision.review_domain
            FROM model_review_heads head
            JOIN model_review_revisions revision ON revision.id = head.revision_id
            WHERE head.model_run_id IN (?, ?)
            ORDER BY revision.id DESC
            """,
            runs,
        ).fetchall()
    by_issue_reviewer: dict[tuple[str, str], dict[str, Any]] = defaultdict(dict)
    for row in rows:
        issue_id = str(row["issue_id"])
        reviewer = str(row["reviewer"] or "").lower()
        if issue_id not in scope_by_issue:
            continue
        by_issue_reviewer[(issue_id, reviewer)][str(row["model_run_id"])] = {
            "status": str(row["status"]),
            "reason": str(row["reason"] or ""),
            "review_domain": str(row["review_domain"] or "model_review"),
        }
    pair = next(
        (
            (issue_id, reviewer, values)
            for (issue_id, reviewer), values in by_issue_reviewer.items()
            if reviewer == author and runs[0] in values and runs[1] in values
        ),
        None,
    )
    if pair is None:
        raise AssertionError("no smoke Run A/Run B current-head pair found")
    issue_id, reviewer, rows_by_run = pair
    encoded_issue = urllib.parse.quote(issue_id, safe="")
    detail_a = api_json(
        base_url,
        f"/api/cases/{encoded_issue}?" + urllib.parse.urlencode({"include_media": "false", "model_run_id": runs[0]}),
    )
    detail_b = api_json(
        base_url,
        f"/api/cases/{encoded_issue}?" + urllib.parse.urlencode({"include_media": "false", "model_run_id": runs[1]}),
    )
    annotations = detail_a.get("annotations") or []
    current_a = next(
        item for item in annotations
        if item.get("review_domain") == "model_review"
        and item.get("model_run_id") == runs[0]
        and str(item.get("author") or "").lower() == reviewer
    )
    current_b = next(
        item for item in annotations
        if item.get("review_domain") == "model_review"
        and item.get("model_run_id") == runs[1]
        and str(item.get("author") or "").lower() == reviewer
    )
    detail_b_runs = {str(item.get("model_run_id") or "") for item in (detail_b.get("annotations") or [])}
    if runs[0] not in detail_b_runs or runs[1] not in detail_b_runs:
        raise AssertionError("case detail did not return the cross-Run audit trail")
    label_a = detail_a.get("label_state") or {}
    label_b = detail_b.get("label_state") or {}
    shared_label_equal = all(
        label_a.get(key) == label_b.get(key)
        for key in ("state", "expected_output", "gt_relation", "source_revision_ids")
    )
    if not shared_label_equal:
        raise AssertionError("shared Label projection differs across selected Runs")

    review_js = (APP_ROOT / "static" / "js" / "review-draft.js").read_text(encoding="utf-8")
    node_script = review_js + r'''
const assert=require('node:assert/strict');
const fs=require('node:fs');
const input=JSON.parse(fs.readFileSync(0,'utf8'));
const state={selectedRunId:input.runA,session:{username:input.reviewer,verified:true}};
const data={annotations:input.annotations};
assert.equal(currentReviewAnnotation(data).model_run_id,input.runA);
assert.equal(currentReviewAnnotation(data).model_review_status,input.statusA);
assert.equal(currentReviewAnnotation(data).note,input.noteA);
state.selectedRunId=input.runB;
assert.equal(currentReviewAnnotation(data).model_run_id,input.runB);
assert.equal(currentReviewAnnotation(data).model_review_status,input.statusB);
assert.equal(currentReviewAnnotation(data).note,input.noteB);
assert.equal(reviewAnnotationsForAllRuns(data).length,input.annotations.length);
'''
    node_input = {
        "annotations": annotations,
        "runA": runs[0],
        "runB": runs[1],
        "reviewer": reviewer,
        "statusA": str(current_a.get("model_review_status") or ""),
        "statusB": str(current_b.get("model_review_status") or ""),
        "noteA": str(current_a.get("note") or ""),
        "noteB": str(current_b.get("note") or ""),
    }
    subprocess.run(
        ["node", "-e", node_script],
        input=json.dumps(node_input, ensure_ascii=False).encode("utf-8"),
        check=True,
        capture_output=True,
        timeout=20,
    )

    with database.connect() as connection:
        comment = connection.execute(
            "SELECT issue_id, model_run_id, body FROM review_comments WHERE body = ? ORDER BY id DESC LIMIT 1",
            ("S3 smoke Run A discussion",),
        ).fetchone()
    if comment is None:
        raise AssertionError("Run-scoped smoke comment not found")
    comment_issue = urllib.parse.quote(str(comment["issue_id"]), safe="")
    comments_a = api_json(
        base_url,
        f"/api/cases/{comment_issue}/comments?" + urllib.parse.urlencode({"model_run_id": runs[0]}),
    )
    comments_b = api_json(
        base_url,
        f"/api/cases/{comment_issue}/comments?" + urllib.parse.urlencode({"model_run_id": runs[1]}),
    )
    bodies_a = {str(item.get("body") or "") for item in comments_a.get("comments", [])}
    bodies_b = {str(item.get("body") or "") for item in comments_b.get("comments", [])}
    if "S3 smoke Run A discussion" not in bodies_a or "S3 smoke Run A discussion" in bodies_b:
        raise AssertionError("Run discussion was not isolated")
    return {
        "run_pair_has_distinct_current_heads": True,
        "detail_returns_full_cross_run_history": True,
        "shared_label_projection_equal_across_runs": True,
        "current_draft_uses_selected_run": True,
        "run_discussion_isolated": True,
        "run_a_status": str(current_a.get("model_review_status") or ""),
        "run_b_status": str(current_b.get("model_review_status") or ""),
    }


def _s2_temporary_snapshot_export_probe(database: Database) -> dict[str, Any]:
    scope = "manual_s3_report_snapshot_probe"
    issue_id = "cn99876543210"
    export_batch_id = ""
    label_case_id = ""
    source_sha = hashlib.sha256(b"manual-s3-report-snapshot-v1").hexdigest()
    membership_sha = _snapshot_membership_sha([issue_id])
    with database.connect() as connection:
        issue_collision = int(connection.execute(
            "SELECT COUNT(*) AS n FROM issues WHERE issue_id=? OR baseline_scope=?",
            (issue_id, scope),
        ).fetchone()["n"])
        state_collision = int(connection.execute(
            "SELECT COUNT(*) AS n FROM gt_sync_state WHERE baseline_scope=?", (scope,)
        ).fetchone()["n"])
    if issue_collision or state_collision:
        raise RuntimeError("temporary S2 probe scope already exists; refusing to overwrite it")
    try:
        database.upsert_issues(
            [{"issue_id": issue_id, "gt_label": "误触发", "gt_source": "smoke_probe"}],
            source="manual_s3_report_probe",
            replace_gt=True,
            baseline_scope=scope,
        )
        with database.connect() as connection:
            connection.execute(
                """
                INSERT INTO gt_sync_state (
                    baseline_scope, status, source_name, source_view_id, source_field,
                    source_sha256, source_row_count, message, error_text
                ) VALUES (?, 'ready', 'smoke_probe', 0, 'probe_gt', ?, 1, 'temporary S2 smoke probe', '')
                """,
                (scope, source_sha),
            )
        first = database.create_gt_snapshot_from_current(
            scope=scope, gt_mode="strict", source_name="smoke_probe", source_view_id=0,
            source_field="probe_gt", created_by="manual_s3_smoke_runner",
            created_by_source="smoke", created_by_verified=False,
            activation_reason="temporary_s2_snapshot_probe",
            expected_member_count=1, expected_membership_sha256=membership_sha,
        )
        reused = database.create_gt_snapshot_from_current(
            scope=scope, gt_mode="strict", source_name="smoke_probe", source_view_id=0,
            source_field="probe_gt", created_by="manual_s3_smoke_runner",
            created_by_source="smoke", created_by_verified=False,
            activation_reason="temporary_s2_snapshot_probe_reuse",
            expected_member_count=1, expected_membership_sha256=membership_sha,
        )
        if reused["id"] != first["id"]:
            raise AssertionError("same GT content did not reuse its snapshot")
        revision = database.create_label_revision(
            issue_id=issue_id, expected_output="正确触发", tags=[], evidence_gaps=[],
            rationale="temporary S2 export reconciliation probe", is_excluded=False,
            author="manual_s3_export_probe", author_source="smoke",
            author_verified=False, expected_previous_revision_id=None,
        )
        label_case_id = str((revision.get("label_case") or {}).get("id") or "")
        ready = [item for item in database.label_gt_candidates([scope]) if item.get("status") == "ready"]
        if not label_case_id or len(ready) != 1:
            raise AssertionError("temporary Label revision did not produce one ready GT candidate")
        preview = database.create_label_gt_export_preview(
            baseline_scopes=[scope], issue_ids=[issue_id],
            created_by="manual_s3_smoke_runner", created_by_source="smoke",
            created_by_verified=False,
        )
        export_batch_id = str(preview["id"])
        database.mark_label_gt_exported(
            batch_id=export_batch_id,
            file_sha256=hashlib.sha256(b"manual-s3-export-probe").hexdigest(),
        )
        not_applied = database.reconcile_label_gt_export_batch(export_batch_id)
        if int(not_applied.get("not_applied_count") or 0) != 1:
            raise AssertionError("GT export did not reconcile as not_applied")
        pinned = database.get_label_gt_export_batch(export_batch_id) or {}
        if str((pinned.get("source_gt_snapshot_ids") or {}).get(scope) or "") != str(first["id"]):
            raise AssertionError("GT export did not retain its original source snapshot")
        with database.connect() as connection:
            connection.execute("UPDATE issues SET gt_label='正确触发' WHERE issue_id=?", (issue_id,))
            connection.execute(
                "UPDATE gt_sync_state SET source_sha256=? WHERE baseline_scope=?",
                (hashlib.sha256(b"manual-s3-report-snapshot-v2").hexdigest(), scope),
            )
        second = database.create_gt_snapshot_from_current(
            scope=scope, gt_mode="strict", source_name="smoke_probe", source_view_id=0,
            source_field="probe_gt", created_by="manual_s3_smoke_runner",
            created_by_source="smoke", created_by_verified=False,
            activation_reason="temporary_s2_snapshot_change",
            expected_member_count=1, expected_membership_sha256=membership_sha,
        )
        matched = database.reconcile_label_gt_export_batch(export_batch_id)
        if int(matched.get("reconcile_counts", {}).get("matched") or 0) != 1:
            raise AssertionError("GT export did not reconcile as matched")
        historical = database.get_gt_snapshot(str(first["id"]), include_items=True, item_limit=5)
        if not historical or str(historical["items"][0].get("gt_label") or "") != "误触发":
            raise AssertionError("historical snapshot facts changed")
        with database.connect() as connection:
            connection.execute("UPDATE issues SET gt_label='无需协助' WHERE issue_id=?", (issue_id,))
            connection.execute(
                "UPDATE gt_sync_state SET source_sha256=? WHERE baseline_scope=?",
                (hashlib.sha256(b"manual-s3-report-snapshot-v3").hexdigest(), scope),
            )
        third = database.create_gt_snapshot_from_current(
            scope=scope, gt_mode="strict", source_name="smoke_probe", source_view_id=0,
            source_field="probe_gt", created_by="manual_s3_smoke_runner",
            created_by_source="smoke", created_by_verified=False,
            activation_reason="temporary_s2_snapshot_change_again",
            expected_member_count=1, expected_membership_sha256=membership_sha,
        )
        changed = database.reconcile_label_gt_export_batch(export_batch_id)
        if int(changed.get("reconcile_counts", {}).get("changed_again") or 0) != 1:
            raise AssertionError("GT export did not reconcile as changed_again")
        if len({str(first["id"]), str(second["id"]), str(third["id"])}) != 3:
            raise AssertionError("changed GT values did not create new immutable snapshots")
        return {
            "same_content_reused_snapshot": True,
            "changed_content_created_new_snapshots": True,
            "historical_snapshot_unchanged": True,
            "temporary_probe_cleanup_verified": True,
            "export_reconcile_statuses": {
                "not_applied": int(not_applied.get("not_applied_count") or 0),
                "matched": int(matched.get("reconcile_counts", {}).get("matched") or 0),
                "changed_again": int(changed.get("reconcile_counts", {}).get("changed_again") or 0),
            },
        }
    finally:
        with database.connect() as connection:
            if export_batch_id:
                connection.execute("DELETE FROM label_gt_export_items WHERE batch_id=?", (export_batch_id,))
                connection.execute("DELETE FROM label_gt_export_source_snapshots WHERE batch_id=?", (export_batch_id,))
                connection.execute("DELETE FROM label_gt_export_batches WHERE id=?", (export_batch_id,))
            if label_case_id:
                connection.execute("DELETE FROM label_resolutions WHERE label_case_id=?", (label_case_id,))
                connection.execute("DELETE FROM label_revisions WHERE label_case_id=?", (label_case_id,))
                connection.execute("DELETE FROM label_cases WHERE id=?", (label_case_id,))
            connection.execute("DELETE FROM gt_snapshot_active WHERE baseline_scope=?", (scope,))
            connection.execute("DELETE FROM gt_snapshot_items WHERE baseline_scope=?", (scope,))
            connection.execute("DELETE FROM gt_snapshots WHERE baseline_scope=?", (scope,))
            connection.execute("DELETE FROM gt_sync_labels WHERE baseline_scope=?", (scope,))
            connection.execute("DELETE FROM gt_sync_state WHERE baseline_scope=?", (scope,))
            connection.execute("DELETE FROM issues WHERE issue_id=?", (issue_id,))
        with database.connect() as connection:
            leftovers = {
                "issue": int(connection.execute("SELECT COUNT(*) AS n FROM issues WHERE issue_id=? OR baseline_scope=?", (issue_id, scope)).fetchone()["n"]),
                "snapshots": int(connection.execute("SELECT COUNT(*) AS n FROM gt_snapshots WHERE baseline_scope=?", (scope,)).fetchone()["n"]),
                "snapshot_items": int(connection.execute("SELECT COUNT(*) AS n FROM gt_snapshot_items WHERE baseline_scope=?", (scope,)).fetchone()["n"]),
                "snapshot_active": int(connection.execute("SELECT COUNT(*) AS n FROM gt_snapshot_active WHERE baseline_scope=?", (scope,)).fetchone()["n"]),
                "sync_state": int(connection.execute("SELECT COUNT(*) AS n FROM gt_sync_state WHERE baseline_scope=?", (scope,)).fetchone()["n"]),
                "label_cases": int(connection.execute("SELECT COUNT(*) AS n FROM label_cases WHERE baseline_scope=?", (scope,)).fetchone()["n"]),
                "export_batch": int(connection.execute("SELECT COUNT(*) AS n FROM label_gt_export_batches WHERE id=?", (export_batch_id,)).fetchone()["n"]) if export_batch_id else 0,
            }
        if any(leftovers.values()):
            raise RuntimeError("temporary S2 probe cleanup failed")


def main() -> int:
    args = parse_args()
    base = urllib.parse.urlparse(args.base_url)
    if base.scheme != "http" or base.hostname not in {"127.0.0.1", "localhost", "::1"} or base.port != 8786:
        raise SystemExit("smoke API base URL must be loopback port 8786")
    manifest_path = Path(args.manifest).expanduser().absolute()
    manifest = _read_check(manifest_path)
    database = Database(read_url_file(Path(args.database_url_file)), postgres_migrations_dir=POSTGRES_MIGRATIONS_DIR, pool_size=4)
    checks: list[dict[str, Any]] = []
    report: dict[str, Any] = {
        "schema": "manual-s1-s3-smoke-report-v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "smoke_database": str(manifest.get("database") or ""),
        "candidate_sha": "",
        "manifest_sha256": hashlib.sha256(manifest_path.read_bytes()).hexdigest(),
        "manifest_path": str(manifest_path),
        "checks": checks,
    }
    try:
        identity = None
        with database.connect() as connection:
            identity = connection.execute(
                "SELECT current_database() AS name, COALESCE(inet_server_addr()::text, '') AS host"
            ).fetchone()
            migration_count = int(connection.execute("SELECT COUNT(*) AS n FROM dashboard_schema_migrations").fetchone()["n"])
        require_database = str(identity["name"] or "")
        require_host = str(identity["host"] or "")
        if not SAFE_DATABASE_RE.fullmatch(require_database) or require_host not in SAFE_HOSTS:
            raise SystemExit("runner refuses non-smoke or non-local database")

        health_path = manifest_path.parent / "health.json"
        if health_path.exists():
            health = _read_check(health_path)
            report["candidate_sha"] = str(health.get("build_commit") or "")

        def sample_integrity() -> dict[str, Any]:
            selected, _scope_map = _selected_ids(manifest)
            builder = _load_smoke_builder()
            active_summary = []
            mismatches = 0
            with database.connect() as connection:
                for item in manifest["scopes"]:
                    ids = [str(row["issue_id"]) for row in connection.execute(
                        "SELECT issue_id FROM issues WHERE baseline_scope=? ORDER BY issue_id",
                        (item["scope"],),
                    ).fetchall()]
                    if len(ids) != int(item["final_count"]) or _snapshot_membership_sha(ids) != item["issue_ids_sha256"]:
                        raise AssertionError("subset membership differs from the manifest")
                    active = database.get_active_gt_snapshot(str(item["scope"]))
                    if not active or str(active.get("id")) != str(item.get("smoke_snapshot_id")):
                        raise AssertionError("active GT snapshot differs from smoke manifest")
                    items = int(connection.execute(
                        "SELECT COUNT(*) AS n FROM gt_snapshot_items WHERE baseline_scope=?",
                        (item["scope"],),
                    ).fetchone()["n"])
                    if items != len(ids):
                        raise AssertionError("GT snapshot item count differs from scope count")
                    active_summary.append({"baseline_id": item["baseline_id"], "members": len(ids), "snapshot_items": items})
                mismatches = int(connection.execute(
                    """
                    SELECT COUNT(*) AS n FROM issues issue
                    JOIN gt_snapshot_active active ON active.baseline_scope=issue.baseline_scope
                    JOIN gt_snapshot_items item ON item.snapshot_id=active.snapshot_id AND item.issue_id=issue.issue_id
                    WHERE issue.baseline_scope IN (?, ?, ?, ?, ?)
                      AND COALESCE(issue.gt_label, '') <> COALESCE(item.gt_label, '')
                    """,
                    tuple(item["scope"] for item in manifest["scopes"]),
                ).fetchone()["n"])
            if len(selected) != int(manifest["selected_issue_count"]) or mismatches:
                raise AssertionError("selected sample count or normalized GT snapshot parity failed")
            return {"selected_issues": len(selected), "scopes": active_summary, "normalized_gt_mismatches": mismatches, "migration_count": migration_count}

        record_check(checks, "S2 subset membership, GT snapshots and normalized overlay", sample_integrity)

        def snapshot_reuse() -> dict[str, Any]:
            results = []
            gt_modes = {
                str(scope): str(mode)
                for _baseline_id, scope, mode in _load_smoke_builder().SCOPES
            }
            for item in manifest["scopes"]:
                before = database.get_active_gt_snapshot(str(item["scope"])) or {}
                reused = database.create_gt_snapshot_from_current(
                    scope=str(item["scope"]), gt_mode=gt_modes[str(item["scope"])],
                    created_by="manual_s3_smoke_runner", created_by_source="smoke",
                    created_by_verified=False, activation_reason="smoke_idempotence_check",
                    expected_member_count=int(item["final_count"]),
                    expected_membership_sha256=str(item["issue_ids_sha256"]),
                )
                after = database.get_active_gt_snapshot(str(item["scope"])) or {}
                if str(reused.get("id")) != str(before.get("id")) or str(after.get("id")) != str(before.get("id")):
                    raise AssertionError("unchanged GT content did not reuse the active snapshot")
                results.append(item["baseline_id"])
            return {"reused_snapshot_count": len(results), "baseline_ids": results}

        record_check(checks, "S2 idempotent snapshot reuse", snapshot_reuse)

        def s1_run_history_and_drafts() -> dict[str, Any]:
            runs = list(manifest["source_run_ids"])
            author = "manual_s3_smoke_gallery_probe"
            with database.connect() as connection:
                rows = connection.execute(
                    """
                    SELECT head.issue_id, head.reviewer, head.model_run_id, revision.status, revision.reason
                    FROM model_review_heads head JOIN model_review_revisions revision ON revision.id=head.revision_id
                    WHERE head.model_run_id IN (?, ?) ORDER BY revision.id DESC
                    """,
                    runs,
                ).fetchall()
            grouped: dict[str, dict[str, Any]] = defaultdict(dict)
            for row in rows:
                if str(row["reviewer"] or "").lower() == author:
                    grouped[str(row["issue_id"])][str(row["model_run_id"])] = {
                        "status": str(row["status"]), "reason": str(row["reason"] or "")
                    }
            pair = next((issue_id, value) for issue_id, value in grouped.items() if runs[0] in value and runs[1] in value)
            issue_id, heads = pair
            encoded = urllib.parse.quote(issue_id, safe="")
            detail_a = api_json(args.base_url, f"/api/cases/{encoded}?" + urllib.parse.urlencode({"include_media": "false", "model_run_id": runs[0]}))
            detail_b = api_json(args.base_url, f"/api/cases/{encoded}?" + urllib.parse.urlencode({"include_media": "false", "model_run_id": runs[1]}))
            history = detail_a.get("annotations") or []
            history_runs = {str(item.get("model_run_id") or "") for item in history}
            if runs[0] not in history_runs or runs[1] not in history_runs:
                raise AssertionError("detail API omitted a Run from the audit history")
            state_a = detail_a.get("label_state") or {}
            state_b = detail_b.get("label_state") or {}
            if any(state_a.get(key) != state_b.get(key) for key in ("state", "expected_output", "gt_relation", "source_revision_ids")):
                raise AssertionError("shared Label state differs between Runs")
            row_a = next(item for item in history if item.get("review_domain") == "model_review" and item.get("model_run_id") == runs[0] and str(item.get("author") or "").lower() == author)
            row_b = next(item for item in history if item.get("review_domain") == "model_review" and item.get("model_run_id") == runs[1] and str(item.get("author") or "").lower() == author)
            node_source = (APP_ROOT / "static" / "js" / "review-draft.js").read_text(encoding="utf-8") + r'''
const assert=require('node:assert/strict');
const fs=require('node:fs');
const input=JSON.parse(fs.readFileSync(0,'utf8'));
const state={selectedRunId:input.runA,session:{username:input.author,verified:true}};
const data={annotations:input.annotations};
assert.equal(currentReviewAnnotation(data).model_run_id,input.runA);
assert.equal(currentReviewAnnotation(data).model_review_status,input.statusA);
assert.equal(currentReviewAnnotation(data).note,input.noteA);
state.selectedRunId=input.runB;
assert.equal(currentReviewAnnotation(data).model_run_id,input.runB);
assert.equal(currentReviewAnnotation(data).model_review_status,input.statusB);
assert.equal(currentReviewAnnotation(data).note,input.noteB);
assert.equal(reviewAnnotationsForAllRuns(data).length,input.annotations.length);
'''
            subprocess.run(
                ["node", "-e", node_source],
                input=json.dumps({"annotations": history, "runA": runs[0], "runB": runs[1], "author": author, "statusA": row_a.get("model_review_status"), "statusB": row_b.get("model_review_status"), "noteA": row_a.get("note") or "", "noteB": row_b.get("note") or ""}, ensure_ascii=False).encode("utf-8"),
                check=True, capture_output=True, timeout=20,
            )
            return {"detail_history_run_count": len(history_runs - {""}), "shared_label_projection_equal": True, "current_draft_run_isolation": True, "run_a_status": heads[runs[0]]["status"], "run_b_status": heads[runs[1]]["status"]}

        record_check(checks, "S1 shared Label state, Run current-head and detail history isolation", s1_run_history_and_drafts)

        def s1_comments() -> dict[str, Any]:
            runs = list(manifest["source_run_ids"])
            with database.connect() as connection:
                row = connection.execute(
                    "SELECT issue_id FROM review_comments WHERE body=? AND model_run_id=? ORDER BY id DESC LIMIT 1",
                    ("S3 smoke Run A discussion", runs[0]),
                ).fetchone()
            if row is None:
                raise AssertionError("smoke Run A comment is missing")
            issue = urllib.parse.quote(str(row["issue_id"]), safe="")
            a = api_json(args.base_url, f"/api/cases/{issue}/comments?" + urllib.parse.urlencode({"model_run_id": runs[0]}))
            b = api_json(args.base_url, f"/api/cases/{issue}/comments?" + urllib.parse.urlencode({"model_run_id": runs[1]}))
            bodies_a = {str(item.get("body") or "") for item in a.get("comments", [])}
            bodies_b = {str(item.get("body") or "") for item in b.get("comments", [])}
            if "S3 smoke Run A discussion" not in bodies_a or "S3 smoke Run A discussion" in bodies_b:
                raise AssertionError("Run comments crossed their Run boundary")
            return {"run_a_comment_visible": True, "run_b_comment_isolated": True}

        record_check(checks, "S1 Run-scoped discussion isolation", s1_comments)

        def s2_temporary_probe() -> dict[str, Any]:
            return _s2_temporary_snapshot_export_probe(database)

        record_check(checks, "S2 changed snapshot history and GT export reconciliation", s2_temporary_probe)

        def s3_api_projection() -> dict[str, Any]:
            runs = list(manifest["source_run_ids"])
            baselines = ",".join(str(item["baseline_id"]) for item in manifest["scopes"])
            reviewer = "manual_s3_smoke_gallery_probe"
            filter_query = urllib.parse.urlencode({"model_run_id": runs[0], "review_status": "in_progress", "annotation_author": reviewer, "baselines": baselines})
            gallery = api_json(args.base_url, "/api/cases?" + filter_query + "&page_size=100")
            analysis = api_json(args.base_url, "/api/review-reason-analysis?" + filter_query)
            if not gallery.get("total") or not analysis.get("total"):
                raise AssertionError("current S3 status filter returned no smoke rows")
            csv_status, _headers, csv_bytes = api_request(
                args.base_url, "GET", "/api/review-reason-analysis/export?" + urllib.parse.urlencode({**{"model_run_id": runs[0], "review_status": "in_progress", "annotation_author": reviewer, "baselines": baselines, "format": "csv"}})
            )
            csv_rows = list(csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
            if csv_status != 200 or not any("in_progress" in row for row in csv_rows[1:]):
                raise AssertionError("CSV export omitted the model-review status")
            xlsx_status, _headers, xlsx_bytes = api_request(
                args.base_url, "GET", "/api/review-reason-analysis/export?" + urllib.parse.urlencode({"model_run_id": runs[0], "review_status": "in_progress", "annotation_author": reviewer, "baselines": baselines, "format": "xlsx"})
            )
            if xlsx_status != 200:
                raise AssertionError("XLSX export failed")
            sheet = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True).active
            if not any("in_progress" in row for row in list(sheet.iter_rows(values_only=True))[1:] if row):
                raise AssertionError("XLSX export omitted the model-review status")
            facets_a = api_json(args.base_url, "/api/model-review-facets?" + urllib.parse.urlencode({"model_run_id": runs[0], "baselines": baselines}))
            facets_b = api_json(args.base_url, "/api/model-review-facets?" + urllib.parse.urlencode({"model_run_id": runs[1], "baselines": baselines}))
            statuses = {str(row["value"]) for row in facets_a.get("statuses", [])} | {str(row["value"]) for row in facets_b.get("statuses", [])}
            if not REQUIRED_MODEL_STATUSES.issubset(statuses):
                raise AssertionError("the four Model Review statuses are not represented in current heads")
            reviewer_rows = list(facets_a.get("reviewers", [])) + list(facets_b.get("reviewers", []))
            smoke_reviewer = next((row for row in reviewer_rows if str(row.get("value") or "").lower() == reviewer), None)
            if not smoke_reviewer or "verified_count" not in smoke_reviewer or "unverified_count" not in smoke_reviewer:
                raise AssertionError("Model Review reviewer facet omitted identity trust counts")
            shadow = api_json(args.base_url, "/api/model-review-shadow-comparison?" + urllib.parse.urlencode({"model_run_id": runs[0], "baselines": baselines}))
            if int(shadow.get("counts", {}).get("matched", 0)) < 1:
                raise AssertionError("the synthetic deterministic legacy backfill did not shadow-match")
            blocked = [row for row in facets_a.get("statuses", []) if row.get("value") == "blocked_by_label"]
            if not blocked:
                raise AssertionError("blocked_by_label facet is missing")
            return {
                "gallery_filter_rows": int(gallery["total"]),
                "analysis_filter_rows": int(analysis["total"]),
                "csv_status_row": True,
                "xlsx_status_row": True,
                "facet_statuses": sorted(statuses),
                "reviewer_facet_has_identity_counts": True,
                "shadow_counts": shadow.get("counts", {}),
            }

        record_check(checks, "S3 statuses, facets, Reason Analysis, CSV/XLSX and shadow projection", s3_api_projection)

        def s3_write_boundaries() -> dict[str, Any]:
            runs = list(manifest["source_run_ids"])
            reviewer = "manual_s3_smoke_gallery_probe"
            with database.connect() as connection:
                row = connection.execute(
                    "SELECT head.issue_id FROM model_review_heads head JOIN model_review_revisions revision ON revision.id=head.revision_id WHERE head.reviewer=? AND head.model_run_id=? AND revision.status='in_progress' ORDER BY revision.id DESC LIMIT 1",
                    (reviewer, runs[0]),
                ).fetchone()
            if row is None:
                raise AssertionError("unassigned Run-bound smoke review not found")
            issue = urllib.parse.quote(str(row["issue_id"]), safe="")
            no_run_review, _headers, _payload = api_request(args.base_url, "POST", f"/api/cases/{issue}/annotations", {"note": "unbound write probe", "author": reviewer})
            no_run_comment, _headers, _payload = api_request(args.base_url, "POST", f"/api/cases/{issue}/comments", {"body": "unbound comment probe", "author": reviewer})
            if no_run_review not in {400, 409} or no_run_comment not in {400, 409}:
                raise AssertionError("unbound legacy write path was not rejected")
            current = database.current_model_review(issue_id=str(row["issue_id"]), model_run_id=runs[0], reviewer=reviewer)
            if not current or current.get("label") or current.get("expected_output") or current.get("tags") or current.get("is_excluded"):
                raise AssertionError("Model Review leaked shared Label fields")
            return {"unbound_review_rejected": no_run_review, "unbound_comment_rejected": no_run_comment, "shared_label_fields_absent": True}

        record_check(checks, "S3 write-domain and unbound legacy boundary", s3_write_boundaries)

        if args.review_attachments_dir:
            record_check(
                checks,
                "S3 binary attachment upload, read-back and cleanup",
                lambda: _binary_attachment_probe(
                    database,
                    manifest,
                    args.base_url,
                    manifest_path,
                    Path(args.review_attachments_dir),
                ),
            )
        else:
            checks.append({"name": "S3 binary attachment upload, read-back and cleanup", "status": "NOT_RUN", "reason": "no isolated smoke attachment directory supplied"})

        def backfill_check() -> dict[str, Any]:
            root = Path(args.backfill_dir).expanduser().absolute() if args.backfill_dir else manifest_path.parent
            if root.resolve() != manifest_path.parent.resolve():
                raise AssertionError("backfill reports must be read from the isolated smoke root")
            initial = _read_check(root / "backfill-preflight.json")
            fixture = _read_check(root / "synthetic-legacy-backfill-fixture.json")
            first_dry = _read_check(root / "backfill-synthetic-dryrun.json")
            first_apply = _read_check(root / "backfill-synthetic-apply.json")
            repeat_dry = _read_check(root / "backfill-synthetic-dryrun2.json")
            repeat_apply = _read_check(root / "backfill-synthetic-apply2.json")
            if initial.get("eligible") or len(initial.get("skipped", {})) != 80:
                raise AssertionError("original mixed legacy preflight did not preserve all 80 ambiguous rows")
            if len(first_dry.get("eligible", [])) != 1 or not fixture.get("synthetic"):
                raise AssertionError("backfill dry-run did not find exactly one marked synthetic row")
            if not first_apply.get("apply") or len(first_apply.get("eligible", [])) != 1:
                raise AssertionError("synthetic backfill apply did not create one record")
            if repeat_dry.get("eligible") or repeat_dry.get("skipped", {}).get(str(fixture["legacy_annotation_id"])) != "already_migrated":
                raise AssertionError("repeated dry-run did not report the synthetic row already migrated")
            if repeat_apply.get("eligible") or repeat_apply.get("skipped", {}).get(str(fixture["legacy_annotation_id"])) != "already_migrated":
                raise AssertionError("repeated apply was not idempotent")
            with database.connect() as connection:
                mapped = connection.execute(
                    "SELECT COUNT(*) AS n FROM model_review_revisions WHERE legacy_annotation_id=? AND status IN ('completed','blocked_by_label')",
                    (int(fixture["legacy_annotation_id"]),),
                ).fetchone()
            if int(mapped["n"] or 0) != 1:
                raise AssertionError("synthetic legacy row does not map to exactly one S3 revision")
            return {"synthetic_eligible": 1, "mixed_legacy_skipped": 80, "repeat_apply_idempotent": True, "synthetic_label_state": fixture.get("label_state", "")}

        record_check(checks, "S3 conservative legacy backfill and idempotence", backfill_check)

        if args.performance_report:
            def performance_check() -> dict[str, Any]:
                performance = _read_check(Path(args.performance_report))
                if int(performance.get("candidate_count", 0)) <= 5000 or int(performance.get("max_batch", 0)) > 400:
                    raise AssertionError("5000+ performance report is missing or batch bound exceeded")
                if int(performance.get("probe_rows_after_cleanup", 1)) != 0:
                    raise AssertionError("performance probe data was not cleaned")
                return {key: performance.get(key) for key in ("candidate_count", "page_items", "batch_count", "max_batch", "page_seconds", "keyset_scan_seconds", "explain")}
            record_check(checks, "5000+ candidate performance probe", performance_check)
        else:
            checks.append({"name": "5000+ candidate performance probe", "status": "NOT_RUN", "reason": "no performance report supplied"})

        def final_integrity() -> dict[str, Any]:
            builder = _load_smoke_builder()
            selected, scope_by_issue = _selected_ids(manifest)
            placeholders, selected_params = builder.selected_table_sql(set(selected))
            with database.connect() as connection:
                orphan_rows = builder.orphan_checks(connection)
                unselected_refs: dict[str, int] = {}
                columns = connection.execute(
                    "SELECT table_name FROM information_schema.columns WHERE table_schema='public' AND column_name='issue_id' ORDER BY table_name"
                ).fetchall()
                for row in columns:
                    table = str(row["table_name"])
                    count = int(connection.execute(
                        f'SELECT COUNT(*) AS n FROM "{table}" WHERE issue_id NOT IN ({placeholders})', selected_params
                    ).fetchone()["n"])
                    if count:
                        unselected_refs[table] = count
                issue_count = int(connection.execute("SELECT COUNT(*) AS n FROM issues").fetchone()["n"])
                normalized_gt_diff = int(connection.execute(
                    """
                    SELECT COUNT(*) AS n FROM issues issue
                    JOIN gt_snapshot_active active ON active.baseline_scope=issue.baseline_scope
                    JOIN gt_snapshot_items item ON item.snapshot_id=active.snapshot_id AND item.issue_id=issue.issue_id
                    WHERE issue.baseline_scope IN (?, ?, ?, ?, ?)
                      AND COALESCE(issue.gt_label,'') <> COALESCE(item.gt_label,'')
                    """,
                    tuple(item["scope"] for item in manifest["scopes"]),
                ).fetchone()["n"])
                artifact_counts = {
                    table: int(connection.execute(f'SELECT COUNT(*) AS n FROM "{table}"').fetchone()["n"])
                    for table in ("review_attachments", "comment_attachments", "model_review_attachments", "review_notifications", "comment_notifications", "intent_comment_notifications", "batch_prediction_jobs", "inference_jobs")
                }
            if orphan_rows or unselected_refs or normalized_gt_diff or issue_count != len(selected):
                raise AssertionError("final smoke database integrity check failed")
            return {"issue_count": issue_count, "foreign_key_orphans": len(orphan_rows), "unselected_reference_tables": len(unselected_refs), "normalized_gt_mismatches": normalized_gt_diff, "artifact_queues_and_attachments": artifact_counts}

        record_check(checks, "final sampled database integrity and side-effect cleanup", final_integrity)
    finally:
        database.close()

    checks.extend(
        [
            {"name": "verified SSO WorkSplit browser acceptance", "status": "NOT_RUN", "reason": "requires an interactive verified SSO session"},
            {"name": "interactive browser screenshot acceptance", "status": "NOT_RUN", "reason": "not performed in this loopback API-only run"},
        ]
    )

    failures = [item for item in checks if item["status"] == "FAIL"]
    report["checks"] = checks
    report["summary"] = {
        "passed": sum(item["status"] == "PASS" for item in checks),
        "failed": len(failures),
        "not_run": sum(item["status"] == "NOT_RUN" for item in checks),
        "selected_issue_count": int(manifest.get("selected_issue_count", 0)),
        "selected_issue_ids_sha256": str(manifest.get("selected_issue_ids_sha256", "")),
    }
    if args.pytest_log:
        pytest_path = Path(args.pytest_log)
        report["pytest_log_path"] = str(pytest_path)
        if pytest_path.is_file():
            lines = [line.strip() for line in pytest_path.read_text(encoding="utf-8", errors="replace").splitlines() if line.strip()]
            report["pytest_summary"] = lines[-1] if lines else ""
    output = Path(args.output).expanduser().absolute()
    output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    temp = output.with_suffix(output.suffix + ".tmp")
    temp.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.chmod(temp, 0o600)
    os.replace(temp, output)
    os.chmod(output, 0o600)
    status = "failed" if failures else "passed_with_gaps" if report["summary"]["not_run"] else "passed"
    print(json.dumps({"status": status, "output": str(output), "summary": report["summary"]}, ensure_ascii=False))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
